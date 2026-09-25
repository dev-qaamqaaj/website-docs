"""OBD AI Solutions: client kickoff deck (16:9), generated with python-pptx.

Uses Playfair Display + Inter (bundled in assets/fonts). Install them on the
presenting machine for the PPTX to look exactly like the PDF export.
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]

BURGUNDY = RGBColor(0x7B, 0x1E, 0x2E)
DEEP = RGBColor(0x4E, 0x0F, 0x1B)
ROSE = RGBColor(0xE7, 0xA3, 0xAE)
INK = RGBColor(0x14, 0x14, 0x14)
TEXT = RGBColor(0x2A, 0x2A, 0x2D)
MUTED = RGBColor(0x6E, 0x6A, 0x6C)
LINE = RGBColor(0xE8, 0xDF, 0xE1)
BLUSH = RGBColor(0xFB, 0xF6, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SERIF = "Playfair Display"
SANS = "Inter"

W, H = Inches(13.333), Inches(7.5)
MX = Inches(0.8)  # side margin


def money(v, cur="PKR"):
    return f"{cur} {v:,.0f}"


# ---------------------------------------------------------------- primitives
def rect(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, runs, size=14, color=TEXT, font=SANS, bold=False, italic=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None, line_spacing=1.1):
    """runs: str | list of paragraphs; a paragraph is str or list of (text, {overrides})."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        if spacing:
            p.space_after = Pt(spacing)
        segs = para if isinstance(para, list) else [(para, {})]
        for seg, o in segs:
            r = p.add_run()
            r.text = seg
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic)
            f.color.rgb = o.get("color", color)
    return tb


def eyebrow(slide, x, y, label, color=BURGUNDY, w=Inches(8)):
    text(slide, x, y, w, Inches(0.3), label.upper(), size=10, color=color, bold=True)


def accent(slide, x, y, color=BURGUNDY):
    rect(slide, x, y, Inches(0.55), Pt(2.2), fill=color)


def logo(slide, x, y, dark=False, size=Inches(0.52)):
    for name in ("obd-logo.png", "obd-logo.jpg"):
        path = ROOT / "assets" / "logo" / name
        if path.exists():
            slide.shapes.add_picture(str(path), x, y, height=size)
            return
    m = rect(slide, x, y, size, size, fill=WHITE if dark else INK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.18)
    tf = m.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "OBD"
    r.font.name, r.font.size, r.font.bold = SERIF, Pt(int(size / Inches(0.52) * 12)), True
    r.font.color.rgb = BURGUNDY if dark else WHITE
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    dot = Emu(int(size * 0.13))
    rect(slide, x + size - dot - Emu(int(size * .12)), y + size - dot - Emu(int(size * .12)), dot, dot,
         fill=RGBColor(0xC2, 0x47, 0x5B), shape=MSO_SHAPE.OVAL)
    text(slide, x + size + Inches(0.14), y + Emu(int(size * .02)), Inches(2.5), Inches(0.3),
         "OBD", size=15, font=SERIF, bold=True, color=WHITE if dark else INK)
    text(slide, x + size + Inches(0.14), y + Emu(int(size * .56)), Inches(2.5), Inches(0.2),
         "A I   S O L U T I O N S", size=7, bold=True, color=ROSE if dark else BURGUNDY)


def chrome(slide, cfg, n, total, section):
    """Standard header/footer for content slides."""
    logo(slide, MX, Inches(0.42), size=Inches(0.42))
    text(slide, W - MX - Inches(4), Inches(0.52), Inches(4), Inches(0.3), section.upper(),
         size=9, color=MUTED, bold=True, align=PP_ALIGN.RIGHT)
    rect(slide, MX, H - Inches(0.62), W - 2 * MX, Pt(0.6), fill=LINE)
    text(slide, MX, H - Inches(0.5), Inches(6), Inches(0.3),
         f"{cfg['company']['name']}  ·  {cfg['project']['code']}", size=8.5, color=MUTED)
    text(slide, W - MX - Inches(2), H - Inches(0.5), Inches(2), Inches(0.3),
         f"{n:02d} / {total:02d}", size=8.5, color=MUTED, align=PP_ALIGN.RIGHT)


def title(slide, kicker, heading, y=Inches(1.25), italic_tail=None, width=Inches(11.5)):
    eyebrow(slide, MX, y, kicker)
    runs = [(heading, {})]
    if italic_tail:
        runs.append((italic_tail, {"italic": True, "color": BURGUNDY}))
    text(slide, MX, y + Inches(0.32), width, Inches(0.9), [runs], size=34, font=SERIF, color=INK)
    accent(slide, MX, y + Inches(1.18))


def card(slide, x, y, w, h, fill=WHITE, border=LINE):
    return rect(slide, x, y, w, h, fill=fill, line=border, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)


def bullets(slide, x, y, w, items, size=13, color=TEXT, mark="✓", mark_color=BURGUNDY, gap=Inches(0.42)):
    for i, it in enumerate(items):
        yy = y + gap * i
        text(slide, x, yy, Inches(0.3), Inches(0.3), mark, size=size, color=mark_color, bold=True)
        text(slide, x + Inches(0.32), yy, w - Inches(0.32), gap, it, size=size, color=color)


# ---------------------------------------------------------------- slides
def s_title(prs, cfg):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, fill=WHITE)
    band = Inches(4.7)
    rect(s, 0, 0, band, H, fill=BURGUNDY)
    rect(s, 0, H - Inches(2.6), band, Inches(2.6), fill=DEEP)
    logo(s, Inches(0.7), Inches(0.7), dark=True)
    text(s, Inches(0.7), H - Inches(2.05), band - Inches(1.2), Inches(1.4),
         [[("Build.  ", {}), ("Launch.  ", {}), ("Care.", {"color": ROSE})]], size=24, font=SERIF, color=WHITE)
    text(s, Inches(0.7), H - Inches(1.1), band - Inches(1.2), Inches(0.4),
         cfg["company"]["tagline"], size=9.5, color=ROSE)

    x = band + Inches(0.9)
    eyebrow(s, x, Inches(1.9), "Project kickoff  ·  Client onboarding")
    text(s, x, Inches(2.3), Inches(7.4), Inches(2.2),
         [[("Your website,", {})], [("launched properly.", {"italic": True, "color": BURGUNDY})]],
         size=46, font=SERIF, color=INK, line_spacing=1.0)
    accent(s, x, Inches(4.2))
    client = cfg["client"]["business"] or cfg["client"]["name"] or "Prepared for our valued client"
    text(s, x, Inches(4.5), Inches(7.2), Inches(0.8),
         "From a finished build to a secure, professional presence on your own domain, with the "
         "infrastructure, email and support behind it.", size=14, color=MUTED, line_spacing=1.3)
    rect(s, x, Inches(6.05), Inches(7.2), Pt(0.6), fill=LINE)
    meta = [("Prepared for", client), ("Presented by", cfg["company"]["founder"]),
            ("Date", cfg["project"]["issue_date"])]
    for i, (k, v) in enumerate(meta):
        xx = x + Inches(2.45) * i
        text(s, xx, Inches(6.2), Inches(2.3), Inches(0.25), k.upper(), size=8, color=MUTED, bold=True)
        text(s, xx, Inches(6.45), Inches(2.3), Inches(0.4), v, size=12, color=INK)


def s_about(prs, cfg, n, t):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(s, cfg, n, t, "Who we are")
    title(s, "01 · OBD AI Solutions", "An engineering studio ", italic_tail="built on AI.")
    text(s, MX, Inches(2.75), Inches(6.2), Inches(1.4),
         "We design and ship AI products, MVPs and production web platforms for founders and growing "
         "businesses, working with clients in the United States, Australia and Pakistan.",
         size=15, color=MUTED, line_spacing=1.35)
    cr = cfg["credentials"]
    stats = [(cr["international_clients"], "US & Australian\nclients served"),
             (cr["local_clients"], "Businesses served\nin Pakistan"),
             ("3", "Countries\nUS · AU · PK")]
    for i, (big, lab) in enumerate(stats):
        x = MX + Inches(2.1) * i
        text(s, x, Inches(4.35), Inches(2), Inches(0.9), big, size=46, font=SERIF, color=BURGUNDY)
        text(s, x, Inches(5.3), Inches(2), Inches(0.7), lab, size=11, color=MUTED, line_spacing=1.2)

    services = [("AI Solutions", "Intelligent automation, LLM assistants and data workflows."),
                ("MVP Development", "Idea to investor-ready product, on foundations that scale."),
                ("Web Engineering", "Fast, secure websites with infrastructure, email and support.")]
    x0, cw = Inches(7.55), Inches(4.98)
    for i, (h, d) in enumerate(services):
        y = Inches(2.7) + Inches(1.28) * i
        card(s, x0, y, cw, Inches(1.1), fill=BLUSH if i != 2 else WHITE, border=BURGUNDY if i == 2 else LINE)
        text(s, x0 + Inches(0.3), y + Inches(0.17), Inches(0.6), Inches(0.4), f"0{i+1}", size=11, bold=True, color=BURGUNDY)
        text(s, x0 + Inches(0.8), y + Inches(0.15), cw - Inches(1), Inches(0.4), h, size=16, font=SERIF, color=INK)
        text(s, x0 + Inches(0.8), y + Inches(0.55), cw - Inches(1), Inches(0.5), d, size=11, color=MUTED)
    text(s, x0 + Inches(0.8), Inches(6.55), cw, Inches(0.3), "← Your project", size=9, bold=True, color=BURGUNDY)


def s_experience(prs, cfg, n, t):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(s, cfg, n, t, "Experience")
    title(s, "02 · Track record", "The teams we've ", italic_tail="built with.")
    names = cfg["credentials"]["notable"]
    cw, gap = Inches(2.78), Inches(0.13)
    for i, nm in enumerate(names):
        x = MX + (cw + gap) * i
        card(s, x, Inches(2.85), cw, Inches(1.55), fill=BLUSH)
        parts = nm.split(" (")
        text(s, x + Inches(0.15), Inches(3.1), cw - Inches(0.3), Inches(0.75), parts[0], size=17, font=SERIF, color=INK,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM)
        if len(parts) > 1:
            text(s, x, Inches(3.95), cw, Inches(0.3), parts[1].rstrip(")").upper(), size=9, bold=True,
                 color=BURGUNDY, align=PP_ALIGN.CENTER)
    text(s, MX, Inches(4.7), Inches(11.7), Inches(0.4),
         f"…plus {cfg['credentials']['local_clients']} local businesses across Pakistan.",
         size=13, color=MUTED, align=PP_ALIGN.CENTER)
    # principles strip
    rect(s, MX, Inches(5.3), W - 2 * MX, Inches(1.2), fill=INK, shape=MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.08
    pr = [("Founder-led", "Omer oversees every launch"), ("Documented", "Scope, fees & support in writing"),
          ("Transparent", "No markups on third-party costs"), ("Yours", "Everything registered to you")]
    for i, (h, d) in enumerate(pr):
        x = MX + Inches(0.4) + Inches(2.93) * i
        text(s, x, Inches(5.5), Inches(2.8), Inches(0.4), h, size=15, font=SERIF, color=WHITE)
        text(s, x, Inches(5.9), Inches(2.8), Inches(0.4), d, size=10.5, color=ROSE)


def s_team(prs, cfg, n, t):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(s, cfg, n, t, "Your team")
    title(s, "03 · Who you'll work with", "Your ", italic_tail="project team.")
    people = [(cfg["company"]["founder"], "Project Lead", [cfg["company"]["founder_title"], cfg["company"]["founder_secondary"]],
               "Owns technical decisions, setup, configuration and launch. Your escalation point for anything critical."),
              (cfg["project"]["account_manager"], "Account Manager", ["Client relations, OBD AI Solutions"],
               "Your day-to-day contact for scheduling, documents, billing and updates.")]
    for i, (nm, role, lines, desc) in enumerate(people):
        x = MX + Inches(5.95) * i
        card(s, x, Inches(2.85), Inches(5.75), Inches(3.4), fill=WHITE, border=BURGUNDY if i == 0 else LINE)
        initials = "".join(w[0] for w in nm.split()[:3] if w[0].isupper())
        c = rect(s, x + Inches(0.4), Inches(3.25), Inches(1.0), Inches(1.0), fill=BURGUNDY if i == 0 else INK, shape=MSO_SHAPE.OVAL)
        tf = c.text_frame
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = initials
        r.font.name, r.font.size, r.font.color.rgb = SERIF, Pt(18), WHITE
        eyebrow(s, x + Inches(1.65), Inches(3.3), role, w=Inches(3.8))
        text(s, x + Inches(1.65), Inches(3.58), Inches(3.9), Inches(0.5), nm, size=22, font=SERIF, color=INK)
        text(s, x + Inches(1.65), Inches(4.08), Inches(3.9), Inches(0.6), lines, size=10.5, color=MUTED, line_spacing=1.2)
        rect(s, x + Inches(0.4), Inches(4.9), Inches(4.95), Pt(0.6), fill=LINE)
        text(s, x + Inches(0.4), Inches(5.05), Inches(4.95), Inches(1), desc, size=12, color=TEXT, line_spacing=1.3)


def s_status(prs, cfg, n, t):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(s, cfg, n, t, "Your project")
    title(s, "04 · Where we are", "Development is done. ", italic_tail="Now we launch.")
    rows = [("Website design & development", "OBD", "Complete", True),
            ("Onboarding & agreement", "Both", "In progress", None),
            ("Domain & hosting purchase (in your name)", "You, guided by us", "Next", False),
            ("Deployment, SSL & security", "OBD", "Scheduled", False),
            ("Professional email & device setup", "OBD", "Scheduled", False),
            ("Go-live, walkthrough & handover", "Both", "Scheduled", False)]
    y0 = Inches(2.85)
    for j, hdr in enumerate(["WORKSTREAM", "OWNER", "STATUS"]):
        text(s, MX + [0, Inches(7.2), Inches(9.9)][j], y0, Inches(3), Inches(0.3), hdr, size=9, bold=True, color=MUTED)
    rect(s, MX, y0 + Inches(0.32), W - 2 * MX, Pt(1.2), fill=INK)
    for i, (w, o, st, done) in enumerate(rows):
        y = y0 + Inches(0.45) + Inches(0.52) * i
        text(s, MX, y, Inches(7), Inches(0.4), w, size=14, color=INK, bold=bool(done))
        text(s, MX + Inches(7.2), y, Inches(2.6), Inches(0.4), o, size=12.5, color=MUTED)
        if done is not None:
            pill = rect(s, MX + Inches(9.9), y - Inches(0.02), Inches(1.45), Inches(0.34),
                        fill=BURGUNDY if done else None, line=None if done else BURGUNDY,
                        shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
            text(s, MX + Inches(9.9), y + Inches(0.03), Inches(1.45), Inches(0.3), st, size=10.5, bold=True,
                 color=WHITE if done else BURGUNDY, align=PP_ALIGN.CENTER)
        else:
            text(s, MX + Inches(9.9), y, Inches(2), Inches(0.4), st, size=12.5, color=MUTED)
        rect(s, MX, y + Inches(0.45), W - 2 * MX, Pt(0.6), fill=LINE)


def s_roadmap(prs, cfg, n, t):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(s, cfg, n, t, "Launch roadmap")
    title(s, "05 · Timeline", "Five working days ", italic_tail="to go live.")
    steps = [("Day 0", "Onboard & sign", "Review pack, onboarding form, agreement, first instalment."),
             ("Day 1–2", "Domain & hosting", "Guided purchase, in your name, on a short call."),
             ("Day 2–4", "Deploy & configure", "Production deploy, DNS, SSL, email and devices."),
             ("Day 5", "Review & go-live", "Final review, launch, walkthrough and handover."),
             ("To " + cfg["project"]["support_until"].split(" ", 1)[1], "Care period", "Support & maintenance, free minor changes in month one.")]
    y_line = Inches(3.35)
    rect(s, MX + Inches(0.15), y_line, W - 2 * MX - Inches(0.3), Pt(1.5), fill=LINE)
    cw = (W - 2 * MX) / 5
    for i, (when, h, d) in enumerate(steps):
        x = MX + cw * i
        last = i == len(steps) - 1
        rect(s, x + Inches(0.05), y_line - Inches(0.14), Inches(0.3), Inches(0.3),
             fill=BURGUNDY if not last else WHITE, line=BURGUNDY, shape=MSO_SHAPE.OVAL)
        text(s, x, Inches(2.75), cw, Inches(0.3), when.upper(), size=10, bold=True, color=BURGUNDY)
        text(s, x, Inches(3.75), cw - Inches(0.25), Inches(0.5), h, size=17, font=SERIF, color=INK)
        text(s, x, Inches(4.5), cw - Inches(0.3), Inches(1.0), d, size=11.5, color=MUTED, line_spacing=1.3)
    card(s, MX, Inches(5.55), W - 2 * MX, Inches(0.85), fill=BLUSH)
    text(s, MX + Inches(0.35), Inches(5.78), W - 2 * MX - Inches(0.7), Inches(0.5),
         [[("Launch quality checklist:  ", {"bold": True, "color": INK}),
           ("HTTPS everywhere  ·  mobile & browser tested  ·  speed reviewed  ·  SPF/DKIM/DMARC email  ·  forms verified  ·  backup handed over", {})]],
         size=12, color=TEXT)


def s_ownership(prs, cfg, n, t):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, fill=INK)
    logo(s, MX, Inches(0.42), dark=True, size=Inches(0.42))
    eyebrow(s, MX, Inches(1.25), "06 · Our ownership promise", color=ROSE)
    text(s, MX, Inches(1.6), Inches(11.5), Inches(1.6),
         [[("Everything that runs your business online ", {}), ("belongs to you.", {"italic": True, "color": ROSE})]],
         size=30, font=SERIF, color=WHITE)
    accent(s, MX, Inches(2.65), color=ROSE)
    items = [("Domain", "Registered in your name with your contact details."),
             ("Hosting", "Account and billing in your name, paid to the provider."),
             ("Email", "Mailboxes on your domain, with passwords that are yours."),
             ("Access", "Delegated for setup only. Revoke it at any time.")]
    cw = (W - 2 * MX) / 4
    for i, (h, d) in enumerate(items):
        x = MX + cw * i
        rect(s, x, Inches(3.1), Pt(1.5), Inches(1.35), fill=BURGUNDY)
        text(s, x + Inches(0.25), Inches(3.1), cw - Inches(0.5), Inches(0.5), h, size=19, font=SERIF, color=WHITE)
        text(s, x + Inches(0.25), Inches(3.6), cw - Inches(0.5), Inches(0.9), d, size=11.5, color=RGBColor(0xC9, 0xC4, 0xC5), line_spacing=1.3)
    rect(s, MX, Inches(4.95), W - 2 * MX, Pt(0.6), fill=RGBColor(0x3A, 0x36, 0x37))
    text(s, MX, Inches(5.15), Inches(4), Inches(0.3), "INDICATIVE THIRD-PARTY COSTS  ·  PAID DIRECTLY BY YOU", size=9, bold=True, color=ROSE)
    for i, tp in enumerate(cfg["pricing"]["third_party_estimates"]):
        x = MX + Inches(3.9) * i
        text(s, x, Inches(5.5), Inches(3.6), Inches(0.6), tp["estimate"], size=26, font=SERIF, color=WHITE)
        text(s, x, Inches(6.1), Inches(3.6), Inches(0.3), tp["item"], size=11, color=RGBColor(0xC9, 0xC4, 0xC5))
    text(s, MX + Inches(7.8), Inches(5.5), Inches(3.9), Inches(1),
         [[("0% commission.", {"bold": True, "color": WHITE})], "We don't resell or mark up domains, hosting or email."],
         size=12, color=RGBColor(0xC9, 0xC4, 0xC5), line_spacing=1.3)
    text(s, W - MX - Inches(2), H - Inches(0.5), Inches(2), Inches(0.3), f"{n:02d} / {t:02d}", size=8.5,
         color=RGBColor(0x8A, 0x85, 0x86), align=PP_ALIGN.RIGHT)


def s_investment(prs, cfg, n, t):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(s, cfg, n, t, "Your investment")
    pr = cfg["pricing"]
    title(s, "07 · Referral pricing", "Below our standard rates, ", italic_tail="by design.")
    text(s, MX, Inches(2.7), Inches(11.5), Inches(0.4),
         f"Because you were referred by {cfg['client']['referred_by']}, preferential rates apply to both development and deployment.",
         size=14, color=MUTED)
    rows = [("Website development", "Completed, built from your documents", pr["development"]["standard"], pr["development"]["offered"])]
    rows += [(f"Deployment · {p['name']}", f"Option {p['id']}" + ("  ·  Recommended" if p["recommended"] else ""),
              pr["deployment_standard"], p["price"]) for p in pr["packages"]]
    y0 = Inches(3.35)
    for j, (hdr, x, al) in enumerate([("SERVICE", MX, PP_ALIGN.LEFT), ("STANDARD", MX + Inches(6.9), PP_ALIGN.RIGHT),
                                       ("YOUR RATE", MX + Inches(9.3), PP_ALIGN.RIGHT)]):
        text(s, x, y0, Inches(2.4) if j else Inches(5), Inches(0.3), hdr, size=9, bold=True, color=MUTED, align=al)
    rect(s, MX, y0 + Inches(0.32), W - 2 * MX, Pt(1.2), fill=INK)
    for i, (a, b, std, off) in enumerate(rows):
        y = y0 + Inches(0.5) + Inches(0.78) * i
        text(s, MX, y, Inches(6.5), Inches(0.4), a, size=16, color=INK, font=SERIF)
        text(s, MX, y + Inches(0.36), Inches(6.5), Inches(0.3), b, size=10.5, color=BURGUNDY if "Recommended" in b else MUTED)
        st = text(s, MX + Inches(6.9), y + Inches(0.05), Inches(2.4), Inches(0.4), money(std, pr["currency"]), size=14,
                  color=RGBColor(0xA2, 0x9C, 0x9E), align=PP_ALIGN.RIGHT)
        st.text_frame.paragraphs[0].runs[0].font._rPr.set("strike", "sngStrike")
        text(s, MX + Inches(9.3), y + Inches(0.02), Inches(2.43), Inches(0.4), money(off, pr["currency"]), size=18,
             bold=True, color=INK, align=PP_ALIGN.RIGHT)
        rect(s, MX, y + Inches(0.7), W - 2 * MX, Pt(0.6), fill=LINE)


def s_packages(prs, cfg, n, t):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(s, cfg, n, t, "Choose your package")
    pr = cfg["pricing"]
    title(s, "08 · Deployment options", "Choose how we ", italic_tail="look after you.")
    cw = Inches(5.75)
    for i, p in enumerate(pr["packages"]):
        x = MX + Inches(5.95) * i
        rec = p["recommended"]
        card(s, x, Inches(2.7), cw, Inches(3.95), fill=WHITE if rec else BLUSH, border=BURGUNDY if rec else LINE)
        if rec:
            rect(s, x, Inches(2.7), cw, Inches(0.09), fill=BURGUNDY)
            pill = rect(s, x + cw - Inches(1.75), Inches(2.95), Inches(1.45), Inches(0.32), fill=BURGUNDY,
                        shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
            text(s, x + cw - Inches(1.75), Inches(3.0), Inches(1.45), Inches(0.3), "RECOMMENDED", size=8.5, bold=True,
                 color=WHITE, align=PP_ALIGN.CENTER)
        eyebrow(s, x + Inches(0.4), Inches(3.0), f"Option {p['id']}", w=Inches(2))
        text(s, x + Inches(0.4), Inches(3.28), Inches(4), Inches(0.5), p["name"], size=22, font=SERIF, color=INK)
        text(s, x + Inches(0.4), Inches(3.8), Inches(4.5), Inches(0.6),
             money(p["price"], pr["currency"]), size=28, font=SERIF, color=BURGUNDY)
        text(s, x + Inches(0.4), Inches(4.35), Inches(5), Inches(0.3),
             f"Total with development: {money(pr['development']['offered'] + p['price'], pr['currency'])}", size=10.5, color=MUTED)
        items = p["includes"]
        for k, it in enumerate(items):
            yy = Inches(4.75) + Inches(0.26) * k
            neg = it.lower().startswith("no ")
            text(s, x + Inches(0.4), yy, Inches(0.3), Inches(0.26), "–" if neg else "✓", size=10.5, bold=True, color=MUTED if neg else BURGUNDY)
            text(s, x + Inches(0.7), yy, cw - Inches(1.0), Inches(0.26), it, size=10.5, color=MUTED if neg else TEXT)


def s_payment(prs, cfg, n, t):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(s, cfg, n, t, "Payment schedule")
    pr = cfg["pricing"]
    dev = pr["development"]["offered"]
    title(s, "09 · When you pay", "Two simple ", italic_tail="instalments.")
    inst = [("Instalment 1", "On signing", "Development fee (100%)\n+ deployment advance (50%)", lambda p: dev + p["price"] / 2),
            ("Instalment 2", "On go-live", "Deployment balance (50%)\nonce your site is live", lambda p: p["price"] / 2)]
    for i, (lab, when, what, fn) in enumerate(inst):
        x = MX + Inches(5.95) * i
        card(s, x, Inches(2.75), Inches(5.75), Inches(2.9), fill=BLUSH if i else WHITE, border=BURGUNDY if not i else LINE)
        text(s, x + Inches(0.4), Inches(3.0), Inches(3), Inches(0.9), f"0{i+1}", size=40, font=SERIF, color=BURGUNDY)
        eyebrow(s, x + Inches(1.5), Inches(3.08), lab, w=Inches(3))
        text(s, x + Inches(1.5), Inches(3.35), Inches(4), Inches(0.5), when, size=20, font=SERIF, color=INK)
        text(s, x + Inches(0.4), Inches(4.05), Inches(5), Inches(0.7), what, size=12, color=MUTED, line_spacing=1.25)
        rect(s, x + Inches(0.4), Inches(4.85), Inches(4.95), Pt(0.6), fill=LINE)
        for j, p in enumerate(pr["packages"]):
            xx = x + Inches(0.4) + Inches(2.5) * j
            text(s, xx, Inches(4.98), Inches(2.4), Inches(0.25), f"OPTION {p['id']} · {p['name'].upper()}", size=8, bold=True, color=MUTED)
            text(s, xx, Inches(5.2), Inches(2.4), Inches(0.4), money(fn(p), pr["currency"]), size=17, bold=True, color=INK)
    text(s, MX, Inches(5.95), W - 2 * MX, Inches(0.5),
         "Bank transfer or mobile wallet  ·  A receipt is issued for every payment  ·  Invoices and a signed agreement support both parties' tax records.",
         size=11.5, color=MUTED, align=PP_ALIGN.CENTER)


def s_support(prs, cfg, n, t):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(s, cfg, n, t, "Care & support")
    title(s, "10 · After launch", "Care that doesn't stop ", italic_tail="at go-live.")
    # left: big date
    card(s, MX, Inches(2.75), Inches(4.3), Inches(3.75), fill=BURGUNDY, border=BURGUNDY)
    text(s, MX + Inches(0.4), Inches(3.05), Inches(3.6), Inches(0.3), "SUPPORTED UNTIL", size=9.5, bold=True, color=ROSE)
    text(s, MX + Inches(0.4), Inches(3.4), Inches(3.6), Inches(1.4), cfg["project"]["support_until"], size=30, font=SERIF, color=WHITE, line_spacing=1.0)
    text(s, MX + Inches(0.4), Inches(4.75), Inches(3.5), Inches(1.6),
         [[(f"+ Free minor UI/content changes for the first {cfg['project']['free_changes_window_days']} days. ", {"bold": True, "color": WHITE})],
          "Something we don't normally include. Complete Care package."],
         size=11.5, color=ROSE, line_spacing=1.3, spacing=4)
    # right: response table
    x = MX + Inches(4.65)
    rows = [("Critical", "Website or email down", "Same business day"),
            ("High", "Key page or form broken", "Within 1 business day"),
            ("Normal", "Minor change or question", "Within 1 business day")]
    text(s, x, Inches(2.75), Inches(2), Inches(0.3), "PRIORITY", size=9, bold=True, color=MUTED)
    text(s, x + Inches(1.6), Inches(2.75), Inches(3), Inches(0.3), "EXAMPLE", size=9, bold=True, color=MUTED)
    text(s, x + Inches(4.9), Inches(2.75), Inches(2.2), Inches(0.3), "FIRST RESPONSE", size=9, bold=True, color=MUTED, align=PP_ALIGN.RIGHT)
    rect(s, x, Inches(3.07), Inches(7.08), Pt(1.2), fill=INK)
    for i, (a, b, c_) in enumerate(rows):
        y = Inches(3.25) + Inches(0.6) * i
        text(s, x, y, Inches(1.5), Inches(0.4), a, size=14, bold=True, color=BURGUNDY if i == 0 else INK)
        text(s, x + Inches(1.6), y, Inches(3.3), Inches(0.4), b, size=13, color=TEXT)
        text(s, x + Inches(4.9), y, Inches(2.18), Inches(0.4), c_, size=13, color=INK, align=PP_ALIGN.RIGHT)
        rect(s, x, y + Inches(0.45), Inches(7.08), Pt(0.6), fill=LINE)
    text(s, x, Inches(5.2), Inches(7.08), Inches(1.3),
         [[("Channels:  ", {"bold": True, "color": INK}), (f"WhatsApp project group  ·  {cfg['project']['account_manager']} (account manager)  ·  {cfg['company']['founder']} (escalations)", {})],
          [("Hours:  ", {"bold": True, "color": INK}), ("Monday–Saturday, 10:00–19:00 PKT", {})]],
         size=12, color=MUTED, line_spacing=1.3, spacing=6)


def s_next(prs, cfg, n, t):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(s, cfg, n, t, "Next steps")
    title(s, "11 · Let's get you live", "Four steps ", italic_tail="from here.")
    steps = [("Review", "This deck and the Welcome & Onboarding Pack."),
             ("Complete", "The Client Onboarding Form: details, domains, email accounts."),
             ("Sign", "Choose a package and sign the Service Agreement. Pay instalment 1."),
             ("Launch", "Book the domain & hosting call. We go live within 5 working days.")]
    cw = (W - 2 * MX - Inches(0.45)) / 4
    for i, (h, d) in enumerate(steps):
        x = MX + (cw + Inches(0.15)) * i
        card(s, x, Inches(2.75), cw, Inches(2.25), fill=BLUSH)
        text(s, x + Inches(0.3), Inches(2.95), Inches(1), Inches(0.8), str(i + 1), size=38, font=SERIF, color=BURGUNDY)
        text(s, x + Inches(0.3), Inches(3.8), cw - Inches(0.5), Inches(0.4), h, size=17, font=SERIF, color=INK)
        text(s, x + Inches(0.3), Inches(4.2), cw - Inches(0.5), Inches(0.8), d, size=11, color=MUTED, line_spacing=1.25)
    text(s, MX, Inches(5.35), Inches(4), Inches(0.3), "YOUR ONBOARDING DOCUMENTS", size=9, bold=True, color=BURGUNDY)
    docs = ["Welcome & Onboarding Pack", "Client Onboarding Form", "Website Service Agreement",
            "Fee & Payment Schedule", "Go-Live Acceptance & Handover"]
    runs = []
    for i, d in enumerate(docs):
        runs += [(f"0{i+1} ", {"bold": True, "color": BURGUNDY}), (d + ("      " if i < len(docs) - 1 else ""), {})]
    text(s, MX, Inches(5.7), W - 2 * MX, Inches(0.4), [runs], size=10, color=INK)


def s_thanks(prs, cfg, n, t):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, fill=BURGUNDY)
    rect(s, 0, H - Inches(1.6), W, Inches(1.6), fill=DEEP)
    logo(s, MX, Inches(0.7), dark=True)
    text(s, MX, Inches(2.3), Inches(11), Inches(1.2), [[("Thank you for ", {}), ("trusting us.", {"italic": True, "color": ROSE})]],
         size=52, font=SERIF, color=WHITE)
    accent(s, MX, Inches(3.65), color=ROSE)
    text(s, MX, Inches(3.95), Inches(9), Inches(1),
         "We're excited to put your business online, and to look after it once it's there.",
         size=17, color=RGBColor(0xF3, 0xDC, 0xE0))
    co = cfg["company"]
    text(s, MX, H - Inches(1.2), Inches(6), Inches(0.4), co["founder"], size=18, font=SERIF, color=WHITE)
    text(s, MX, H - Inches(0.8), Inches(8), Inches(0.4), f"{co['founder_title']}  ·  {co['founder_secondary']}", size=10.5, color=ROSE)
    contact = "  ·  ".join(v for v in (co["email"], co["phone"], co["website"]) if v)
    if contact:
        text(s, W - MX - Inches(5), H - Inches(1.0), Inches(5), Inches(0.4), contact, size=11, color=WHITE, align=PP_ALIGN.RIGHT)


def build(cfg, out_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    body = [s_about, s_experience, s_team, s_status, s_roadmap, s_ownership, s_investment,
            s_packages, s_payment, s_support, s_next]
    total = len(body) + 2
    s_title(prs, cfg)
    for i, fn in enumerate(body, start=2):
        fn(prs, cfg, i, total)
    s_thanks(prs, cfg, total, total)
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    return out_path
